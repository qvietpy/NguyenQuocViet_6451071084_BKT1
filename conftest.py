import platform
import shutil
import time
from datetime import datetime
from pathlib import Path

import pytest
import selenium
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from e2e.base.base_test import BaseTest
from e2e.pages.login_page import LoginPage


REPORT_DIR = Path("reports")
SCREENSHOT_DIR = Path("screenshots")
TEST_RESULTS = []
SESSION_STARTED = time.perf_counter()


@pytest.fixture
def driver():
    browser = BaseTest.create_driver()
    yield browser
    browser.quit()


@pytest.fixture
def login_page(driver):
    return LoginPage(driver)


def _case_data(item):
    marker = item.get_closest_marker("case")
    return marker.kwargs if marker else {}


def _safe_screenshot(item, report):
    driver = item.funcargs.get("driver")
    if not driver or not report.failed:
        return ""

    SCREENSHOT_DIR.mkdir(exist_ok=True)
    case_id = _case_data(item).get("id", item.name)
    path = SCREENSHOT_DIR / f"{case_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
    driver.save_screenshot(str(path))
    return str(path)


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    outcome = yield
    report = outcome.get_result()

    if report.when != "call":
        return

    data = _case_data(item)
    status = "PASS" if report.passed else "SKIP" if report.skipped else "FAIL"
    screenshot = _safe_screenshot(item, report)

    if report.passed:
        actual = f"Assertion passed: {data.get('expected', item.name)}"
    elif report.skipped:
        actual = str(report.longrepr)
    else:
        actual = str(report.longrepr)
        if len(actual) > 1500:
            actual = actual[:1500] + "..."

    TEST_RESULTS.append(
        {
            "id": data.get("id", item.name),
            "title": data.get("title", item.name),
            "username": data.get("username", ""),
            "password": data.get("password", ""),
            "expected": data.get("expected", ""),
            "actual": actual,
            "status": status,
            "duration": round(report.duration, 3),
            "screenshot": screenshot,
        }
    )


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(sheet):
    for column in sheet.columns:
        column_letter = get_column_letter(column[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def _write_results_sheet(workbook):
    sheet = workbook.active
    sheet.title = "Test Results"
    sheet.append(
        [
            "Test Case ID",
            "Title",
            "Username",
            "Password",
            "Expected Result",
            "Actual Result",
            "Status",
            "Duration (seconds)",
            "Screenshot",
        ]
    )
    _style_header(sheet)

    fills = {
        "PASS": PatternFill("solid", fgColor="C6EFCE"),
        "FAIL": PatternFill("solid", fgColor="FFC7CE"),
        "SKIP": PatternFill("solid", fgColor="FFEB9C"),
    }

    for result in TEST_RESULTS:
        sheet.append(
            [
                result["id"],
                result["title"],
                result["username"],
                result["password"],
                result["expected"],
                result["actual"],
                result["status"],
                result["duration"],
                result["screenshot"],
            ]
        )
        sheet.cell(sheet.max_row, 7).fill = fills.get(result["status"], fills["FAIL"])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _auto_width(sheet)


def _write_summary_sheet(workbook, elapsed):
    total = len(TEST_RESULTS)
    passed = sum(1 for result in TEST_RESULTS if result["status"] == "PASS")
    failed = sum(1 for result in TEST_RESULTS if result["status"] == "FAIL")
    skipped = sum(1 for result in TEST_RESULTS if result["status"] == "SKIP")
    pass_rate = round((passed / total * 100), 2) if total else 0

    sheet = workbook.create_sheet("Summary")
    for row in [
        ("Project", "NguyenQuocViet_6451071084_BKT1"),
        ("Target URL", LoginPage.URL),
        ("Execution Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Test Cases", total),
        ("Passed", passed),
        ("Failed", failed),
        ("Skipped", skipped),
        ("Pass Rate (%)", pass_rate),
        ("Execution Time (seconds)", round(elapsed, 3)),
        ("Python version", platform.python_version()),
        ("Selenium version", selenium.__version__),
    ]:
        sheet.append(row)
    _style_header(sheet)
    _auto_width(sheet)


def pytest_sessionfinish(session, exitstatus):
    REPORT_DIR.mkdir(exist_ok=True)
    elapsed = time.perf_counter() - SESSION_STARTED
    workbook = Workbook()
    _write_results_sheet(workbook)
    _write_summary_sheet(workbook, elapsed)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = REPORT_DIR / f"test_report_{timestamp}.xlsx"
    latest_path = REPORT_DIR / "latest_test_report.xlsx"
    workbook.save(report_path)
    shutil.copyfile(report_path, latest_path)
    print(f"\nExcel report generated: {report_path}")
    print(f"Latest report updated: {latest_path}")
